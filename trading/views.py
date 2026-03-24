from django.shortcuts import render
from django.http import JsonResponse, HttpResponse
from django.contrib.auth.decorators import login_required
from django.views.decorators.http import require_POST
from django.db import transaction as db_transaction
from decimal import Decimal, InvalidOperation
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet
import io
import json 
import csv
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from stocks.models import Stock, PriceBar
from . models import Portfolio, execute_market_order, Trade, Position, PendingOrder, PortfolioSnapshot


@login_required
@require_POST
def place_market_order(request):
    try:
        data = json.loads(request.body)
        symbol = data.get('symbol', '').upper().strip()
        trade_type = data.get('trade_type', '').upper().strip()
        raw_quantity = data.get('quantity', 0)

        if trade_type not in ('BUY', 'SELL'):
            return JsonResponse({"success":False, 'error':'Invalid trade type'}, status=400)
        
        try:
            quantity = Decimal(str(raw_quantity))
        except InvalidOperation:
            return JsonResponse({'success':False, 'error':'Invalid quantity'}, status=400)
        
        if quantity <= 0:
            return JsonResponse({'success':False, 'error':'Quantity must be greater than 0'}, status=400)
        
        
        stock = Stock.objects.get(symbol = symbol)
        portfolio = request.user.portfolio

        trade = execute_market_order(portfolio, stock, trade_type, quantity)

        portfolio.refresh_from_db()

        return JsonResponse({
            'success':True,
            'message':f'{trade_type} {quantity} {symbol} @ {trade.price}',
            'new_cash':str(portfolio.cash),
            'trade_price':str(trade.price),
            'total_value':str(trade.total_value),
        })
        
    except Stock.DoesNotExist:
        return JsonResponse({'success':False, 'error':f'Stock not found'}, status=400)
    except ValueError as e:
        return JsonResponse({'success':False, 'error':str(e)}, status=400)
    except Exception as e:
        return JsonResponse({'success':False, 'error':'Something went wrong'}, status=500)
    

@login_required
def portfolio_dashboard(request):
    portfolio = request.user.portfolio
    positions = portfolio.positions.filter(quantity__gt=0).select_related('stock')

    enriched_positions = []
    total_holdings_value = Decimal('0')

    for pos in positions:
        latest = PriceBar.objects.filter(stock=pos.stock).order_by('-timestamp').first()

        if latest:
            current_price = latest.close
        else:
            current_price = pos.avg_cost
        
        current_value = current_price * pos.quantity
        cost_basis = pos.avg_cost * pos.quantity
        pnl = current_value - cost_basis
        pnl_pct = (pnl / cost_basis * 100) if  cost_basis > 0 else Decimal('0')

        total_holdings_value += current_value

        enriched_positions.append({
            'stock':pos.stock,
            'quantity':pos.quantity,
            'avg_cost':pos.avg_cost,
            'current_price':current_price,
            'current_value':current_value,
            'pnl':pnl,
            'pnl_pct':pnl_pct,
        })
    
    total_portfolio_value = portfolio.cash + total_holdings_value
    total_pnl = total_portfolio_value - Decimal('100000')

    pending_orders = PendingOrder.objects.filter(
        portfolio=portfolio, 
        status='PENDING'
        ).select_related('stock').order_by('-created_at')
    

    recent_trades = portfolio.trades.select_related('stock').order_by('-executed_at')[:10]

    return render(request, 'trading/dashboard.html', {
        'portfolio':portfolio,
        'positions':enriched_positions,
        'pending_orders':pending_orders,
        'recent_trades':recent_trades,
        'total_portfolio_value':total_portfolio_value,
        'total_holdings_value':total_holdings_value,
        'total_pnl':total_pnl
    })

@login_required
@require_POST
def place_limit_order(request):
    try:
        data = json.loads(request.body)
        symbol = data.get('symbol', '').upper().strip()
        order_type = data.get('order_type', '').upper().strip()
        raw_quantity = data.get('quantity', 0)
        raw_trigger_price = data.get('trigger_price', 0)

        valid_types = ('LIMIT_BUY', 'LIMIT_SELL', 'STOP_LOSS', 'TAKE_PROFIT')
        if order_type not in valid_types:
            return JsonResponse({'success':False, 'error':'Invalid order type'}, status=400)
        
        try:
            quantity = Decimal(str(raw_quantity))
            trigger_price = Decimal(str(raw_trigger_price))

        except InvalidOperation:
            return JsonResponse({'success':False, 'error':'Invalid quantity or price'}, status=400)
        
        if quantity <= 0 or trigger_price <= 0:
            return JsonResponse({'success':False, 'error':'Quantity and price must be greater than 0'}, status=400)
        
        stock = Stock.objects.get(symbol=symbol)
        portfolio = request.user.portfolio

        order = PendingOrder.objects.create(
            portfolio=portfolio,
            stock=stock,
            order_type=order_type,
            quantity=quantity,
            trigger_price=trigger_price,
        )
        return JsonResponse({
            'success':True,
            'message':f"{order_type} order placed - {quantity} {symbol} @ {trigger_price}",
            'order_id':order.id,
        })
    except Stock.DoesNotExist:
        return JsonResponse({'success':False, 'error':'Stock not found'}, status=404)
    except Exception as e:
        return JsonResponse({'success':False, 'error':'Something went wrong'}, status=500)
    

@login_required
@require_POST
def cancel_limit_order(request, order_id):
    try:
        order = PendingOrder.objects.get(
            id=order_id,
            portfolio=request.user.portfolio,
            status='PENDING'
        )
        order.status = 'CANCELLED'
        order.save()
        return JsonResponse({'success':True, 'message':'Order cancelled'})
    except PendingOrder.DoesNotExist:
        return JsonResponse({'success':False, 'error':'Order not found'}, status=404)
    

@login_required
@require_POST
def place_stop_order(request):
    try:
        data = json.loads(request.body)
        symbol = data.get('symbol', '').upper().strip()
        order_type = data.get('order_type', '').upper().strip()
        raw_quantity = data.get('quantity',0)
        raw_trigger_price = data.get('trigger_price', 0)

        if order_type not in ('STOP_LOSS', 'TAKE_PROFIT'):
            return JsonResponse({'success':False, 'error':'Invalid order type'}, status=400)
        
        try:
            quantity = Decimal(str(raw_quantity))
            trigger_price = Decimal(str(raw_quantity))
        except InvalidOperation:
            return JsonResponse({'success':False,'error':'Invalid quantity or price'},status=400)
        
        if quantity <=0 or trigger_price <= 0:
            return JsonResponse({'success':False, 'error':'Quantity and price must be greater than 0'},status=400)
        
        stock = Stock.objects.get(symbol=symbol)
        portfolio = request.user.portfolio

        latest = PriceBar.objects.filter(stock=stock).order_by('-timestamp').first()
        if not latest:
            return JsonResponse({"success":False, "message":"No price data available"}, status=400)
        
        current_price = latest.close
        if order_type == 'STOP_LOSS' and trigger_price >= current_price:
            return JsonResponse({
                'success':False,
                'error':f'stop-loss price {trigger_price} must be below current price {current_price}'
            }, status=400)
        
        if order_type == 'TAKE_PROFIT' and trigger_price <= current_price:
            return JsonResponse({
                'success':False,
                'error':f'take-profit price {trigger_price}, must be above current price {current_price}'
            }, status=400)
        
        try:
            position = Position.objects.get(portfolio=portfolio, stock=stock)
            if position.quantity < quantity:
                return JsonResponse({
                    'success':False,
                    'error':f"You only have {position.quantity} shares, cannot set order for {quantity}"
                }, status=400)
        
        except Position.DoesNotExist:
            return JsonResponse({
                'success':False,
                'message':f"You do not hold any {symbol} shares"
            }, status=400)
            
        order = PendingOrder.objects.create(
            portfolio=portfolio,
            stock=stock,
            order_type=order_type,
            quantity=quantity,
            trigger_price=trigger_price,
        )
        return JsonResponse({
            'success':True,
            'message':f"{order_type} set for {quantity} {symbol} @ {trigger_price}",
            'order_id':order.id
        })
        
    except Stock.DoesNotExist:
            return JsonResponse({'success':False, 'error':'Stock not found'}, status=400)
    except Exception as e:
            return JsonResponse({'success':False, 'error':'Something went wrong'}, status=500)
        

@login_required
def pnl_history_json(request):
    snapshots = PortfolioSnapshot.objects.filter(
        portfolio=request.user.portfolio
    ).order_by('timestamp')

    if not snapshots.exists():
        return JsonResponse({
            'labels':['Start'],
            'values':[100000.00],
        })
    
    return JsonResponse({
        'labels': [s.timestamp.strftime('%b %d %H:%M') for s in snapshots],
        'values': [float(s.total_value) for s in snapshots],
    })


@login_required
def trade_history(request):
    trades = portfolio_trades = Trade.objects.filter(
        portfolio = request.user.portfolio
    ).select_related('stock').order_by('-executed_at')

    symbol_filter = request.GET.get('symbol', '').upper().strip()
    if symbol_filter:
        trades = trades.filter(stock__symbol=symbol_filter)

    date_from = request.GET.get('date_from', '').strip()
    date_to = request.GET.get('date_to', '').strip()

    if date_from:
        trades = trades.filter(executed_at__date__gte=date_from)
    if date_to:
        trades = trades.filter(executed_at__date__lte=date_to)

    export = request.GET.get('export','')

    if export == 'csv':
        response = HttpResponse(content_type='text/csv')
        response['Content-Disposition'] = 'attachment; filename="trade_history.csv"'
        writer = csv.writer(response)
        writer.writerow(['Date', 'Symbol', 'Type', 'Quantity', 'Price', 'Total'])
        for t in trades:
            writer.writerow([
                t.executed_at.strftime('%Y-%m-%d %H:%M'),
                t.stock.symbol,
                t.trade_type,
                float(t.quantity),
                float(t.price),
                float(t.quantity * t.price)
            ])
        return response
    
    if export == 'xlsx':
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Trade History'

        headers = ['Date', 'Symbol', 'Type', 'Quantity', 'Price', 'Total']
        ws.append(headers)

        for cell in ws[1]:
            cell.font = Font(bold=True, color='FFFFFF') 
            cell.fill = PatternFill(start_color='1a1a2e', end_color='1a1a2e', fill_type='solid')
            cell.alignment = Alignment(horizontal='center')

        for t in trades:
            ws.append([
                t.executed_at.strftime('%Y-%m-%d %H:%M'),
                t.stock.symbol,
                t.trade_type,
                float(t.quantity),
                float(t.price),
                float(t.quantity * t.price)
            ])
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            ws.column_dimensions[col[0].column_letter].width = max_len + 4

        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = 'attachment; filename="trade_history.xlsx"'
        wb.save(response)
        return response
    
    if export == 'pdf':
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=landscape(A4))
        styles = getSampleStyleSheet()
        elements = []

        # Title
        elements.append(Paragraph('Trade History', styles['Title']))
        elements.append(Spacer(1, 12))

        # Table data
        data = [['Date', 'Symbol', 'Type', 'Quantity', 'Price', 'Total']]
        for t in trades:
            data.append([
                t.executed_at.strftime('%Y-%m-%d %H:%M'),
                t.stock.symbol,
                t.trade_type,
                str(float(t.quantity)),
                f'${float(t.price):.2f}',
                f'${float(t.quantity * t.price):.2f}',
            ])

        table = Table(data, repeatRows=1)
        table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1a1a2e')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 11),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f5f5f5')]),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
            ('FONTSIZE', (0, 1), (-1, -1), 9),
            ('TOPPADDING', (0, 0), (-1, -1), 6),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
        ]))
        elements.append(table)
        doc.build(elements)

        buffer.seek(0)
        response = HttpResponse(buffer, content_type='application/pdf')
        response['Content-Disposition'] = 'attachment; filename="trade_history.pdf"'
        return response

    total_trades = trades.count()
    total_bought = trades.filter(trade_type='BUY').count()
    total_sold = trades.filter(trade_type='SELL').count()

    return render(request, 'trading/trade_history.html',{
        'trades':trades,
        'symbol_filter':symbol_filter,
        'date_from':date_from,
        'date_to':date_to,
        'total_trades':total_trades,
        'total_bought':total_bought,
        'total_sold':total_sold,
    })

@login_required
@require_POST
def reset_portfolio(request):
    with db_transaction.atomic():
        old_portfolio = request.user.portfolio
        old_portfolio.is_active = False
        old_portfolio.save()

        new_portfolio = Portfolio.objects.create(
            user = request.user,
            cash = 100000,
            is_active = True
        )
    return JsonResponse({
        'success':True,
        'message':'Portfolio reset. Starting fresh with 100,000.',
        'redirect':'/trading/dashboard/',
    })